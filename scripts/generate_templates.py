"""Generate survey_template.csv and survey_template.xlsx."""

from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from mei.io.schema import load_schema

schema = load_schema()

# 1. Canonical column list (35 columns)
canonical_columns = [col.name for col in schema.columns]

# Create survey_template.csv (headers only)
csv_path = Path("data/templates/survey_template.csv")
csv_path.parent.mkdir(parents=True, exist_ok=True)
empty_df = pd.DataFrame(columns=canonical_columns)
empty_df.to_csv(csv_path, index=False, encoding="utf-8")
print(f"Created {csv_path} with {len(canonical_columns)} headers.")

# 2. Build Data Dictionary rows
dict_rows = []
for col in schema.columns:
    options_str = ""
    if col.text_labels:
        options_str = "; ".join([f"{k} = {v}" for k, v in col.text_labels.items()])
    elif col.allowed_codes:
        options_str = f"Allowed codes: {col.allowed_codes}"

    dict_rows.append({
        "Column Name": col.name,
        "Role": col.role,
        "Period": f"P{col.period}" if col.period else "-",
        "Indicator / Item": col.indicator or col.item_id or "-",
        "Data Type": col.type,
        "Description": col.description or col.statement or "-",
        "Allowed Codes & Text Options": options_str or "-"
    })

dict_df = pd.DataFrame(dict_rows)

# 3. Write Excel with two sheets: 'Survey Data' and 'Data Dictionary'
xlsx_path = Path("data/templates/survey_template.xlsx")
with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
    empty_df.to_excel(writer, sheet_name="Survey Data", index=False)
    dict_df.to_excel(writer, sheet_name="Data Dictionary", index=False)

    # Style the workbook with openpyxl
    wb = writer.book
    ws1 = wb["Survey Data"]
    ws2 = wb["Data Dictionary"]

    # Header styling
    header_fill = PatternFill(start_color="1F4E46", end_color="1F4E46", fill_type="solid") # Primary forest
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for ws in [ws1, ws2]:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Auto-adjust column widths
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

print(f"Created {xlsx_path} with sheets ['Survey Data', 'Data Dictionary'].")
